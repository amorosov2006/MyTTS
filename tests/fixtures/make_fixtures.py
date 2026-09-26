#!/usr/bin/env python3
"""Generate small test books in multiple formats from shared content."""

import json
import os
import subprocess
import base64
import zipfile
from pathlib import Path
from io import BytesIO
from datetime import datetime

# Third-party imports
import pymupdf as fitz  # PyMuPDF
from reportlab.lib.pagesizes import A5
from reportlab.lib.units import cm, mm
from reportlab.pdfgen import canvas
from docx import Document
from PIL import Image, ImageDraw, ImageFont

try:
    from ebooklib import epub
    HAS_EPUB = True
except:
    HAS_EPUB = False

# Book content definitions
RUSSIAN_BOOK = {
    "slug": "quiet_station_ru",
    "title": "Тихая станция",
    "author": "Иван Петров",
    "lang": "ru",
    "chapters": [
        {
            "title": "Глава 1. Туман",
            "key_sentence": "На улице висел густой туман, затрудняющий проход даже в полдень.",
            "paragraphs": [
                "На улице висел густой туман, затрудняющий проход даже в полдень. Станция казалась особенно одинокой в такую погоду. Редкие путники проходили мимо закрытых лавок.",
                "Смотритель станции сидел в своей конторе, читая старую газету 1891 г. Мир казался ему большим и пугающим. У здания стояло ровно 3 дома, разбросанных по краям площади.",
                "Вечером город погрузился в глубокую тишину. Свечи в окнах горели тускло и покачивались на ветру. Никто не приходил и не уходил через станцию."
            ]
        },
        {
            "title": "Глава 2. Письмо",
            "key_sentence": "Письмо пришло в самый неожиданный момент, когда смотритель уже не ждал никаких новостей.",
            "paragraphs": [
                "Письмо пришло в самый неожиданный момент, когда смотритель уже не ждал никаких новостей.[1] Конверт был помят и грязен, словно долго путешествовал по дорогам. Адрес написан красивым почерком, хотя и с ошибками.",
                "Смотритель открыл письмо дрожащими руками. «— Вы опоздали, — сказал он себе вслух, прочитав первые строки.» Его сердце забилось учащённо. Было ещё слишком много неясного.",
                "Он понимал, что теперь всё должно измениться. Прошлое возвращалось неожиданно. Будущее казалось неопределённым и тревожным."
            ]
        },
        {
            "title": "Глава 3. Мост",
            "key_sentence": "Мост через реку был единственным путём в новую жизнь, но смотритель боялся его пересечь.",
            "paragraphs": [
                "Мост через реку был единственным путём в новую жизнь, но смотритель боялся его пересечь. Деревянные доски скрипели под ногами и качались на ветру. Вода внизу казалась бесконечно глубокой и чёрной.",
                "Он стоял посередине моста, глядя на темный горизонт. Позади остались годы сомнений и неудач. Впереди ждало что-то неизвестное, но это уже не пугало его как раньше.",
                "Наконец он переступил через последние доски. Противоположный берег встретил его молчаливо и спокойно. И он понял, что большая часть его жизни уже позади, но новая только начинается."
            ]
        }
    ],
    "footnote_text": "Примечание автора: это вымышленная история."
}

ENGLISH_BOOK = {
    "slug": "quiet_station_en",
    "title": "The Quiet Station",
    "author": "John Peters",
    "lang": "en",
    "chapters": [
        {
            "title": "Chapter 1. Fog",
            "key_sentence": "A thick fog hung over the street, making passage difficult even in broad daylight.",
            "paragraphs": [
                "A thick fog hung over the street, making passage difficult even in broad daylight. The station seemed especially lonely in such weather. Few travelers passed by the closed shops.",
                "The station master sat in his office, reading an old newspaper from 1891. The world seemed vast and frightening to him. There were exactly three houses standing near the edges of the square.",
                "As evening fell, the city sank into deep silence. Candles in windows burned dimly and flickered in the wind. No one came or left through the station."
            ]
        },
        {
            "title": "Chapter 2. The Letter",
            "key_sentence": "The letter arrived at the most unexpected moment, when the station master had long stopped hoping for any news.",
            "paragraphs": [
                "The letter arrived at the most unexpected moment, when the station master had long stopped hoping for any news.[1] The envelope was wrinkled and dirty, as if it had traveled long roads. The address was written in a beautiful hand, though with some errors.",
                "The station master opened the letter with trembling hands. \"You were late,\" said Dr. Smith in a voice that seemed to come from the past. His heart began to race. There was still too much unclear.",
                "He understood that everything must change now. The past returned unexpectedly. The future seemed uncertain and alarming, like a fog that would not lift."
            ]
        },
        {
            "title": "Chapter 3. The Bridge",
            "key_sentence": "The bridge across the river was the only path to a new life, but the station master feared to cross it.",
            "paragraphs": [
                "The bridge across the river was the only path to a new life, but the station master feared to cross it. The wooden boards creaked beneath his feet and swayed in the wind. The water below seemed infinitely deep and black, worth perhaps $2,500 in his dreams.",
                "He stood in the middle of the bridge, gazing at the dark horizon. Behind him lay years of doubt and failure. Ahead lay something unknown, but it no longer frightened him as it once did.",
                "At last he stepped across the final boards. The opposite shore greeted him silently and peacefully. And he understood that the greater part of his life was behind him, but a new one was just beginning."
            ]
        }
    ],
    "footnote_text": "Author's note: this is a fictional story."
}


def generate_cover_image(title: str, author: str, output_path: str) -> bytes:
    """Generate a simple cover image (400x600 PNG)."""
    img = Image.new('RGB', (400, 600), color=(70, 130, 180))  # steel blue
    draw = ImageDraw.Draw(img)

    # Try to use a better font; fall back to default if not available
    try:
        title_font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 32)
        author_font = ImageFont.truetype("/System/Library/Fonts/Helvetica.ttc", 20)
    except:
        title_font = ImageFont.load_default()
        author_font = ImageFont.load_default()

    # Draw title
    draw.text((20, 250), title, fill=(255, 255, 255), font=title_font)
    # Draw author
    draw.text((20, 350), f"by {author}", fill=(200, 200, 200), font=author_font)

    img.save(output_path)
    with open(output_path, 'rb') as f:
        return f.read()


def create_epub(book: dict, output_path: str) -> None:
    """Create an EPUB file using ebooklib."""
    if not HAS_EPUB:
        print(f"    ✗ EPUB skipped (ebooklib not available)")
        return

    try:
        book_epub = epub.EpubBook()

        # Metadata
        book_epub.set_identifier(f"mytest_{book['slug']}")
        book_epub.set_title(book['title'])
        book_epub.set_language(book['lang'])
        book_epub.add_author(book['author'])

        # Cover image
        cover_path = output_path.replace('.epub', '_cover.png')
        cover_data = generate_cover_image(book['title'], book['author'], cover_path)
        cover_image = epub.EpubImage()
        cover_image.uid = 'cover'
        cover_image.file_name = 'cover.png'
        cover_image.content = cover_data
        book_epub.add_item(cover_image)

        # Front matter
        c0 = epub.EpubHtml()
        c0.uid = 'cover_page'
        c0.file_name = '00_cover.xhtml'
        c0.content = f'''<html><head><title>Cover</title></head><body>
    <h1>{book['title']}</h1>
    <p>by {book['author']}</p>
    <p style="margin-top: 2em;">Copyright and licensing information.</p>
    </body></html>'''
        book_epub.add_item(c0)

        chapters = []

        for ch_idx, chapter in enumerate(book['chapters']):
            c = epub.EpubHtml()
            c.uid = f'chapter_{ch_idx + 1}'
            c.file_name = f'chapter_{ch_idx + 1}.xhtml'

            # Build content with footnote in chapter 2
            paragraphs = ''.join([f'<p>{p}</p>' for p in chapter['paragraphs']])

            # Add footnote to chapter 2
            if ch_idx == 1:
                paragraphs = paragraphs.replace(
                    '.[1]',
                    '.<a epub:type="noteref" href="#n1" id="n1_ref">1</a>'
                )
                # Add footnote section
                footnote_html = f'''<aside epub:type="footnote" id="n1">
            <p><a epub:type="backlink" href="#n1_ref">1</a>. {book['footnote_text']}</p>
            </aside>'''
                paragraphs = paragraphs + footnote_html

            c.content = f'''<html><head><title>{chapter['title']}</title></head><body>
        <h1>{chapter['title']}</h1>
        {paragraphs}
        </body></html>'''

            book_epub.add_item(c)
            chapters.append(c)

        # Table of contents
        book_epub.toc = tuple(chapters)

        # Add navigation
        book_epub.add_item(epub.EpubNcx())
        book_epub.add_item(epub.EpubNav())

        book_epub.spine = ['nav', c0] + chapters

        epub.write_epub(output_path, book_epub, {})
    except Exception as e:
        print(f"    ✗ EPUB generation failed: {e}")


def create_fb2(book: dict, output_path: str) -> None:
    """Create an FB2 XML file."""
    cover_path = output_path.replace('.fb2', '_cover.png')
    cover_data = generate_cover_image(book['title'], book['author'], cover_path)
    cover_b64 = base64.b64encode(cover_data).decode('ascii')

    # Build chapters
    chapters_xml = ''
    for ch_idx, chapter in enumerate(book['chapters']):
        paragraphs_xml = ''.join([f'<p>{p}</p>' for p in chapter['paragraphs']])

        # Add footnote reference in chapter 2
        if ch_idx == 1:
            paragraphs_xml = paragraphs_xml.replace(
                '.[1]',
                '.<a l:href="#n1" type="note">1</a>'
            )

        chapters_xml += f'''
        <section>
            <title><p>{chapter['title']}</p></title>
            {paragraphs_xml}
        </section>
        '''

    # Build footnotes section
    footnotes_xml = f'''
    <section>
        <title><p>Примечания</p></title>
        <p><a id="n1" type="note">[1]</a> {book['footnote_text']}</p>
    </section>
    '''

    # Full FB2 document
    fb2_content = f'''<?xml version="1.0" encoding="UTF-8"?>
<FictionBook xmlns="http://www.gribuser.ru/xml/fictionbook/2.0" xmlns:l="http://www.w3.org/1999/xlink">
    <description>
        <title-info>
            <genre>fiction</genre>
            <book-title>{book['title']}</book-title>
            <author>
                <first-name>{"Иван" if book['lang'] == 'ru' else "John"}</first-name>
                <last-name>{"Петров" if book['lang'] == 'ru' else "Peters"}</last-name>
            </author>
            <lang>{book['lang']}</lang>
        </title-info>
    </description>
    <body>
        <coverpage>
            <image l:href="#cover"/>
        </coverpage>
        {chapters_xml}
    </body>
    <body name="notes">
        {footnotes_xml}
    </body>
    <binary id="cover" content-type="image/png">
        {cover_b64}
    </binary>
</FictionBook>
'''

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(fb2_content)


def create_fb2_zip(fb2_path: str) -> None:
    """Create a ZIP file containing the FB2."""
    zip_path = fb2_path.replace('.fb2', '.fb2.zip')
    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as z:
        z.write(fb2_path, arcname=Path(fb2_path).name)


def create_docx(book: dict, output_path: str) -> None:
    """Create a DOCX file."""
    doc = Document()

    # Set core properties
    doc.core_properties.title = book['title']
    doc.core_properties.author = book['author']

    # Title
    title_para = doc.add_paragraph(book['title'])
    title_para.style = 'Heading 1'

    # Author
    doc.add_paragraph(f"by {book['author']}")
    doc.add_paragraph()  # blank line

    footnotes_list = []

    for ch_idx, chapter in enumerate(book['chapters']):
        # Chapter heading
        ch_para = doc.add_paragraph(chapter['title'])
        ch_para.style = 'Heading 1'

        # Paragraphs
        for para_text in chapter['paragraphs']:
            # Add footnote reference in chapter 2
            if ch_idx == 1 and '[1]' in para_text:
                footnotes_list.append(book['footnote_text'])

            p = doc.add_paragraph(para_text)
            p.style = 'Normal'

        doc.add_paragraph()  # blank line between chapters

    # Add footnotes section
    if footnotes_list:
        doc.add_paragraph()
        notes_heading = doc.add_paragraph("Примечания" if book['lang'] == 'ru' else "Notes")
        notes_heading.style = 'Heading 1'
        for note_text in footnotes_list:
            p = doc.add_paragraph(note_text)
            p.style = 'Normal'

    doc.save(output_path)


def find_font_with_cyrillic() -> str:
    """Find a TrueType font that supports Cyrillic."""
    candidates = [
        "/System/Library/Fonts/Supplemental/Arial Unicode.ttf",
        "/Library/Fonts/Arial Unicode.ttf",
        "/System/Library/Fonts/Helvetica.ttc",
        "/System/Library/Fonts/Menlo.ttc"
    ]
    for path in candidates:
        if os.path.exists(path):
            return path
    # Fallback: look in /System/Library/Fonts for any TTF
    import glob
    fonts = glob.glob("/System/Library/Fonts/*.ttf") + glob.glob("/System/Library/Fonts/*.ttc")
    if fonts:
        return fonts[0]
    return None


def create_pdf(book: dict, output_path: str) -> None:
    """Text PDF: Unicode font, running header, page-number footer, chapter bookmarks,
    words hyphenated across line breaks, footnote at the bottom of chapter 2."""
    from reportlab.lib.pagesizes import A6
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    font_name = "ArialUnicode"
    if font_name not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont(font_name, find_font_with_cyrillic()))
    c = canvas.Canvas(output_path, pagesize=A6)
    c.setTitle(book['title'])
    c.setAuthor(book['author'])
    w, h = A6
    margin = 1.0 * cm
    width = w - 2 * margin
    size = 10
    page_no = [1]

    def decorate():
        c.setFont(font_name, 7)
        c.drawString(margin, h - 0.6 * cm, book['title'])
        c.line(margin, h - 0.75 * cm, w - margin, h - 0.75 * cm)
        c.drawCentredString(w / 2, 0.5 * cm, str(page_no[0]))
        c.setFont(font_name, size)

    def new_page():
        c.showPage()
        page_no[0] += 1
        decorate()
        return h - 1.3 * cm

    def wrap(text):
        lines, cur = [], ""
        for word in text.split():
            trial = f"{cur} {word}".strip()
            if c.stringWidth(trial, font_name, size) <= width:
                cur = trial
                continue
            if len(word) >= 8 and word.isalpha():  # hyphenate long words across the break
                cut = len(word) // 2
                head = f"{cur} {word[:cut]}-".strip()
                if c.stringWidth(head, font_name, size) <= width:
                    lines.append(head)
                    cur = word[cut:]
                    continue
            lines.append(cur)
            cur = word
        return lines + ([cur] if cur else [])

    decorate()
    for ch_idx, chapter in enumerate(book['chapters']):
        if ch_idx > 0:
            y = new_page()
        else:
            y = h - 1.3 * cm
        key = f"ch{ch_idx}"
        c.bookmarkPage(key)
        c.addOutlineEntry(chapter['title'], key, level=0)
        c.setFont(font_name, 13)
        c.drawString(margin, y, chapter['title'])
        y -= 0.8 * cm
        c.setFont(font_name, size)
        for para in chapter['paragraphs']:
            for i, line in enumerate(wrap(para)):
                if y < 1.2 * cm:
                    y = new_page()
                c.drawString(margin + (0.4 * cm if i == 0 else 0), y, line)
                y -= 0.42 * cm
            y -= 0.15 * cm
        if ch_idx == 1:
            c.setFont(font_name, 7)
            c.line(margin, 1.45 * cm, margin + 2 * cm, 1.45 * cm)
            c.drawString(margin, 1.1 * cm, f"1 {book['footnote_text']}")
            c.setFont(font_name, size)
    c.save()


def create_scanned_pdf(pdf_path: str, output_path: str) -> None:
    """Image-only PDF (no text layer) of the first 2 pages at 200 dpi, for OCR tests."""
    src = fitz.open(pdf_path)
    out = fitz.open()
    for page in list(src)[:2]:
        pix = page.get_pixmap(dpi=200, colorspace=fitz.csGRAY)
        dst = out.new_page(width=page.rect.width, height=page.rect.height)
        dst.insert_image(dst.rect, stream=pix.tobytes("png"))
    out.save(output_path, deflate=True)


def create_txt(book: dict, output_path: str, encoding: str = 'utf-8', wrap: bool = False) -> None:
    """Create a plain text file."""
    lines = [
        book['title'],
        book['author'],
        '',
    ]

    for chapter in book['chapters']:
        lines.append('')
        lines.append(chapter['title'])
        lines.append('')

        for para in chapter['paragraphs']:
            if wrap:
                # Hard-wrap at ~70 columns
                wrapped = []
                words = para.split()
                current_line = ''
                for word in words:
                    if len(current_line) + len(word) + 1 > 70:
                        wrapped.append(current_line)
                        current_line = word
                    else:
                        current_line = current_line + (' ' if current_line else '') + word
                if current_line:
                    wrapped.append(current_line)
                lines.extend(wrapped)
            else:
                lines.append(para)
            lines.append('')

    content = '\n'.join(lines)
    with open(output_path, 'w', encoding=encoding) as f:
        f.write(content)


def create_markdown(book: dict, output_path: str) -> None:
    """Create a Markdown file."""
    lines = [
        f"# {book['title']}",
        '',
        f"**Author:** {book['author']}",
        '',
    ]

    for chapter in book['chapters']:
        lines.append(f"## {chapter['title']}")
        lines.append('')

        for para in chapter['paragraphs']:
            lines.append(para)
            lines.append('')

        lines.append('```')
        lines.append('code block that should be skipped')
        lines.append('```')
        lines.append('')

    # Add a link
    lines.append("[More information](https://example.com)")
    lines.append('')

    content = '\n'.join(lines)
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(content)


def create_html(book: dict, output_path: str) -> None:
    """Create an HTML file."""
    html = f'''<!DOCTYPE html>
<html lang="{book['lang']}">
<head>
    <meta charset="UTF-8">
    <title>{book['title']}</title>
</head>
<body>
    <h1>{book['title']}</h1>
    <p><strong>Author:</strong> {book['author']}</p>
'''

    for chapter in book['chapters']:
        html += f'    <h2>{chapter["title"]}</h2>\n'
        for para in chapter['paragraphs']:
            html += f'    <p>{para}</p>\n'

    html += '</body>\n</html>'

    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(html)


def convert_with_textutil(html_path: str, format_type: str) -> str:
    """Convert HTML to other formats using textutil (macOS)."""
    output_path = html_path.replace('.html', f'.{format_type}')
    try:
        subprocess.run(
            ['textutil', '-convert', format_type, html_path, '-output', output_path],
            check=True,
            capture_output=True,
            timeout=30
        )
        return output_path
    except subprocess.CalledProcessError as e:
        return None
    except Exception as e:
        return None


def verify_epub(epub_path: str) -> bool:
    """Verify EPUB file."""
    try:
        book = epub.read_epub(epub_path)
        items = list(book.get_items())
        return len(items) > 0
    except Exception:
        return False


def verify_pdf(pdf_path: str) -> bool:
    """Verify PDF file."""
    try:
        doc = fitz.open(pdf_path)
        result = len(doc) > 0
        doc.close()
        return result
    except Exception:
        return False


def verify_docx(docx_path: str) -> bool:
    """Verify DOCX file."""
    try:
        doc = Document(docx_path)
        return len(doc.paragraphs) > 0
    except Exception:
        return False


def generate_expected_json(output_dir: str) -> None:
    """Generate expected.json with test metadata."""
    expected = {
        "quiet_station_ru": {
            "title": "Тихая станция",
            "author": "Иван Петров",
            "lang": "ru",
            "chapter_titles": ["Глава 1. Туман", "Глава 2. Письмо", "Глава 3. Мост"],
            "key_sentences": [
                "На улице висел густой туман, затрудняющий проход даже в полдень.",
                "Письмо пришло в самый неожиданный момент, когда смотритель уже не ждал никаких новостей.",
                "Мост через реку был единственным путём в новую жизнь, но смотритель боялся его пересечь."
            ],
            "footnote_text": "Примечание автора: это вымышленная история.",
            "formats": {
                "epub": ["title", "author", "lang", "chapters", "footnote"],
                "fb2": ["title", "author", "lang", "chapters", "footnote"],
                "docx": ["title", "author", "lang", "chapters", "footnote"],
                "pdf": ["title", "lang", "chapters"],
                "txt": ["lang", "chapters"],
                "md": ["lang", "chapters"],
                "html": ["title", "lang", "chapters"]
            }
        },
        "quiet_station_en": {
            "title": "The Quiet Station",
            "author": "John Peters",
            "lang": "en",
            "chapter_titles": ["Chapter 1. Fog", "Chapter 2. The Letter", "Chapter 3. The Bridge"],
            "key_sentences": [
                "A thick fog hung over the street, making passage difficult even in broad daylight.",
                "The letter arrived at the most unexpected moment, when the station master had long stopped hoping for any news.",
                "The bridge across the river was the only path to a new life, but the station master feared to cross it."
            ],
            "footnote_text": "Author's note: this is a fictional story.",
            "formats": {
                "epub": ["title", "author", "lang", "chapters", "footnote"],
                "fb2": ["title", "author", "lang", "chapters", "footnote"],
                "docx": ["title", "author", "lang", "chapters", "footnote"],
                "pdf": ["title", "lang", "chapters"],
                "txt": ["lang", "chapters"],
                "md": ["lang", "chapters"],
                "html": ["title", "lang", "chapters"]
            }
        }
    }

    output_path = os.path.join(output_dir, 'expected.json')
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(expected, f, ensure_ascii=False, indent=2)


def main():
    """Generate all test books."""
    output_dir = Path(__file__).parent / 'books'
    output_dir.mkdir(parents=True, exist_ok=True)

    # Clean cover images if they exist
    for book in [RUSSIAN_BOOK, ENGLISH_BOOK]:
        cover_path = str(output_dir / f"{book['slug']}_cover.png")
        if os.path.exists(cover_path):
            os.remove(cover_path)

    print("Generating test books...")
    results = {}

    for book in [RUSSIAN_BOOK, ENGLISH_BOOK]:
        slug = book['slug']
        print(f"\n  {slug}:")
        results[slug] = {}

        # EPUB
        epub_path = str(output_dir / f"{slug}.epub")
        create_epub(book, epub_path)
        if os.path.exists(epub_path):
            size = os.path.getsize(epub_path)
            results[slug]['epub'] = (epub_path, size, verify_epub(epub_path))
            print(f"    ✓ EPUB ({size} bytes)")

        # FB2
        fb2_path = str(output_dir / f"{slug}.fb2")
        create_fb2(book, fb2_path)
        size = os.path.getsize(fb2_path)
        results[slug]['fb2'] = (fb2_path, size, size > 0)
        print(f"    ✓ FB2 ({size} bytes)")

        # FB2 ZIP
        fb2_zip_path = str(output_dir / f"{slug}.fb2.zip")
        create_fb2_zip(fb2_path)
        size = os.path.getsize(fb2_zip_path)
        results[slug]['fb2_zip'] = (fb2_zip_path, size, size > 0)
        print(f"    ✓ FB2 ZIP ({size} bytes)")

        # DOCX
        docx_path = str(output_dir / f"{slug}.docx")
        create_docx(book, docx_path)
        size = os.path.getsize(docx_path)
        results[slug]['docx'] = (docx_path, size, verify_docx(docx_path))
        print(f"    ✓ DOCX ({size} bytes)")

        # PDF
        pdf_path = str(output_dir / f"{slug}.pdf")
        create_pdf(book, pdf_path)
        size = os.path.getsize(pdf_path)
        results[slug]['pdf'] = (pdf_path, size, verify_pdf(pdf_path))
        print(f"    ✓ PDF ({size} bytes)")

        # Scanned PDF
        scanned_pdf_path = str(output_dir / f"{slug}_scanned.pdf")
        create_scanned_pdf(pdf_path, scanned_pdf_path)
        if os.path.exists(scanned_pdf_path):
            size = os.path.getsize(scanned_pdf_path)
            results[slug]['scanned_pdf'] = (scanned_pdf_path, size, verify_pdf(scanned_pdf_path))
            print(f"    ✓ Scanned PDF ({size} bytes)")

        # TXT
        txt_path = str(output_dir / f"{slug}.txt")
        create_txt(book, txt_path, encoding='utf-8', wrap=False)
        size = os.path.getsize(txt_path)
        results[slug]['txt'] = (txt_path, size, size > 0)
        print(f"    ✓ TXT ({size} bytes)")

        # Additional TXT variants for Russian
        if book['lang'] == 'ru':
            # CP1251 variant
            cp1251_path = str(output_dir / f"{slug}_cp1251.txt")
            create_txt(book, cp1251_path, encoding='cp1251', wrap=False)
            size = os.path.getsize(cp1251_path)
            results[slug]['txt_cp1251'] = (cp1251_path, size, size > 0)
            print(f"    ✓ TXT (CP1251) ({size} bytes)")

            # Wrapped variant
            wrapped_path = str(output_dir / f"{slug}_wrapped.txt")
            create_txt(book, wrapped_path, encoding='utf-8', wrap=True)
            size = os.path.getsize(wrapped_path)
            results[slug]['txt_wrapped'] = (wrapped_path, size, size > 0)
            print(f"    ✓ TXT (wrapped) ({size} bytes)")

        # Markdown
        md_path = str(output_dir / f"{slug}.md")
        create_markdown(book, md_path)
        size = os.path.getsize(md_path)
        results[slug]['md'] = (md_path, size, size > 0)
        print(f"    ✓ Markdown ({size} bytes)")

        # HTML
        html_path = str(output_dir / f"{slug}.html")
        create_html(book, html_path)
        size = os.path.getsize(html_path)
        results[slug]['html'] = (html_path, size, size > 0)
        print(f"    ✓ HTML ({size} bytes)")

        # Converted formats with textutil
        for fmt in ['rtf', 'odt', 'doc']:
            result = convert_with_textutil(html_path, fmt)
            if result and os.path.exists(result):
                size = os.path.getsize(result)
                results[slug][fmt] = (result, size, size > 0)
                print(f"    ✓ {fmt.upper()} ({size} bytes)")

    # Generate expected.json
    generate_expected_json(str(output_dir))
    print("\n  ✓ expected.json")

    # Summary
    print("\n" + "=" * 60)
    print("SUMMARY")
    print("=" * 60)

    total_size = 0
    file_count = 0

    for slug in [RUSSIAN_BOOK['slug'], ENGLISH_BOOK['slug']]:
        print(f"\n{slug}:")
        for fmt, (path, size, valid) in results[slug].items():
            status = "✓" if valid else "?"
            print(f"  {status} {fmt:12} {size:8} bytes  {Path(path).name}")
            total_size += size
            file_count += 1

    print(f"\n{'=' * 60}")
    print(f"Total: {file_count} files, {total_size / 1024 / 1024:.2f} MB")
    print(f"{'=' * 60}")

    if total_size > 3 * 1024 * 1024:
        print(f"Warning: total size exceeds 3 MB limit!")
    else:
        print(f"✓ Within 3 MB limit ({(3 * 1024 * 1024 - total_size) / 1024:.1f} KB remaining)")


if __name__ == '__main__':
    main()
